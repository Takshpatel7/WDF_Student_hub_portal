from pathlib import Path
import importlib.util

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Student_Hub_Practical_File_1_to_6_All_Code.docx"


def load_module(name, path):
    spec = importlib.util.spec_from_file_location(name, path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


base = load_module("base_builder", ROOT / "build_practical_file.py")
p6 = load_module("practical6_builder", ROOT / "build_practical6.py")


def title_page(doc):
    p = doc.add_paragraph(style="Title")
    base.style_paragraph(p, before=100, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Student Hub Portal")
    base.set_run_font(r, size=25, bold=True, color="000000")

    p = doc.add_paragraph()
    base.style_paragraph(p, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Practical File 1 to 6")
    base.set_run_font(r, size=20, bold=True, color="1F4E79")

    p = doc.add_paragraph()
    base.style_paragraph(p, after=14, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("HTML CSS JavaScript JSON and supporting source code")
    base.set_run_font(r, size=12, color="555555")

    p = doc.add_paragraph()
    base.style_paragraph(p, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Questions answers analysis interpretation and complete code")
    base.set_run_font(r, size=11, color="555555")
    doc.add_page_break()


def add_practical6(doc):
    base.add_heading(doc, "Practical No. 6 - Rendering External JSON Data using Fetch API, Search and Filter")
    base.add_body(doc, "Question", "Display event lists, student profiles, notices, or FAQs from an external JSON file using Fetch API. Implement dynamic rendering, search, filter, sorting, and pagination.")
    base.add_subheading(doc, "Key Questions and Answers")
    questions = [
        ("How is JSON fetched, parsed, and rendered?", "The fetch function requests events.json and returns a response. The response is converted to JavaScript objects with response.json(), and the objects are inserted into the page as event cards."),
        ("Which array methods are used for search, filter, sort, and pagination?", "filter() selects matching events, sort() orders events by date, map() creates display markup, and slice() selects the records for the current page."),
        ("How are loading and error states handled?", "The page initially shows a loading message. If the request fails or the response is not successful, catch() displays an error message and clears the event list."),
        ("How is modularity maintained in JavaScript files?", "The code is divided into small functions such as loadEvents, addCategories, applyFilters, renderEvents, renderPagination, and formatDate. Each function has one clear responsibility."),
    ]
    for question, answer in questions:
        base.add_body(doc, "Q", question)
        base.add_body(doc, "Answer", answer)
    base.add_subheading(doc, "Analysis")
    base.add_body(doc, "A", "The application loads event records from events.json instead of writing every event directly in HTML. Search, category filtering, date sorting, and pagination are applied to the JavaScript array before the visible cards are updated.")
    base.add_subheading(doc, "Interpretation")
    base.add_body(doc, "I", "This approach is useful for modern web applications because the same page can display changing data without rewriting the HTML. In a real project, the JSON file can later be replaced by an API endpoint.")
    base.add_subheading(doc, "Required Practical 6 Code")
    base.add_body(doc, "Testing Note", "Run the files through a local server such as Live Server. Fetch may be blocked when the HTML file is opened directly with a file path in the browser.")
    for label, code in [
        ("HTML Code - events-fetch.html", p6.HTML),
        ("JSON Code - events.json", p6.JSON_DATA),
        ("JavaScript Code - practical6.js", p6.JS),
        ("CSS Code - practical6.css", p6.CSS),
    ]:
        base.add_code_heading(doc, label)
        base.add_code_block(doc, code)


def add_code_appendix(doc):
    doc.add_page_break()
    base.add_heading(doc, "Complete Source Code Appendix")
    base.add_body(doc, "Included Files", "The following appendix contains every code file currently present in the final folder: HTML, CSS, JavaScript, JSON, and Python source files. Generated Word/PDF files and image assets are not included as code.")
    extensions = {".html", ".css", ".js", ".json", ".py"}
    files = sorted((p for p in ROOT.iterdir() if p.is_file() and p.suffix.lower() in extensions), key=lambda p: (p.suffix.lower(), p.name.lower()))
    for path in files:
        base.add_code_heading(doc, f"Source Code - {path.name}")
        base.add_code_block(doc, path.read_text(encoding="utf-8"))


def main():
    doc = Document()
    base.configure_document(doc)
    title_page(doc)
    for item_index, item in enumerate(base.practical_data_upto_5()):
        if item_index:
            doc.add_page_break()
        base.add_practical(doc, *item)
    doc.add_page_break()
    add_practical6(doc)
    add_code_appendix(doc)
    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
