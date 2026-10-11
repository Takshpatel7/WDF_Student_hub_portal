from pathlib import Path
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "Student_Hub_Practical_6_Fetch_API_Search_Filter.docx"


def font(run, name="Calibri", size=11, bold=False, color="000000"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def para(p, before=0, after=4, line=1.15, align=None):
    f = p.paragraph_format
    f.space_before = Pt(before)
    f.space_after = Pt(after)
    f.line_spacing = line
    if align is not None:
        p.alignment = align


def heading(doc, text, level=1):
    p = doc.add_paragraph()
    para(p, before=12 if level == 1 else 8, after=5, line=1.05)
    r = p.add_run(text)
    font(r, size=16 if level == 1 else 12.5, bold=True, color="1F4E79")


def labeled(doc, label, text):
    p = doc.add_paragraph()
    para(p, after=4)
    r = p.add_run(label + ": ")
    font(r, bold=True, color="1F4E79")
    r = p.add_run(text)
    font(r)


def code_color(line):
    stripped = line.strip()
    if not stripped:
        return "666666"
    if stripped.startswith("//") or stripped.startswith("/*") or stripped.startswith("*") or stripped.startswith("<!--"):
        return "008000"
    if stripped.startswith("{") or stripped.startswith("}") or stripped.startswith('"'):
        return "A31515"
    if re.search(r"<(/?)(html|head|body|main|section|div|h1|h2|p|input|button|script|link|select|option|form|label|title)", stripped):
        return "0000CC"
    if re.search(r"\b(function|const|let|return|async|await|if|else|for|from|import|class)\b", stripped):
        return "7F0055"
    if ":" in stripped and stripped.endswith(";"):
        return "006666"
    return "333333"


def code_block(doc, title, code):
    p = doc.add_paragraph()
    para(p, before=8, after=3, line=1.0)
    r = p.add_run(title)
    font(r, size=11, bold=True, color="1F4E79")
    for line in code.rstrip().splitlines():
        p = doc.add_paragraph()
        para(p, after=0, line=1.0)
        p.paragraph_format.left_indent = Inches(0.15)
        r = p.add_run(line)
        font(r, name="Courier New", size=9, color=code_color(line))


HTML = '''<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <title>Student Hub Events</title>
  <link rel="stylesheet" href="practical6.css">
</head>
<body>
  <main class="container">
    <h1>Student Hub Events</h1>
    <p>Events are loaded from an external JSON file using the Fetch API.</p>

    <section class="controls" aria-label="Event controls">
      <label for="searchInput">Search</label>
      <input id="searchInput" type="search" placeholder="Search by event or venue">
      <label for="categoryFilter">Category</label>
      <select id="categoryFilter">
        <option value="all">All categories</option>
      </select>
      <label for="sortSelect">Sort by date</label>
      <select id="sortSelect">
        <option value="asc">Oldest first</option>
        <option value="desc">Newest first</option>
      </select>
    </section>

    <p id="status" class="status" role="status">Loading events...</p>
    <section id="eventList" class="event-grid" aria-live="polite"></section>
    <nav id="pagination" class="pagination" aria-label="Event pages"></nav>
  </main>
  <script src="practical6.js"></script>
</body>
</html>'''


JSON_DATA = '''[
  {"id": 1, "title": "SPOURAL 2026", "date": "2026-10-25", "venue": "College Main Ground", "category": "Sports"},
  {"id": 2, "title": "Cyber Security Workshop", "date": "2026-11-10", "venue": "Auditorium B", "category": "Technical"},
  {"id": 3, "title": "Career Guidance Seminar", "date": "2026-11-22", "venue": "Seminar Hall", "category": "Career"},
  {"id": 4, "title": "Student Hackathon", "date": "2026-12-05", "venue": "Innovation Lab", "category": "Technical"},
  {"id": 5, "title": "Annual Cultural Night", "date": "2026-12-18", "venue": "Open Air Theatre", "category": "Cultural"},
  {"id": 6, "title": "Resume Building Session", "date": "2027-01-08", "venue": "Placement Cell", "category": "Career"}
]'''


JS = '''const state = {
  events: [],
  filtered: [],
  page: 1,
  pageSize: 4
};

const list = document.getElementById("eventList");
const status = document.getElementById("status");
const searchInput = document.getElementById("searchInput");
const categoryFilter = document.getElementById("categoryFilter");
const sortSelect = document.getElementById("sortSelect");
const pagination = document.getElementById("pagination");

async function loadEvents() {
  try {
    const response = await fetch("events.json");
    if (!response.ok) throw new Error("Unable to load event data");
    state.events = await response.json();
    addCategories();
    applyFilters();
  } catch (error) {
    status.textContent = error.message;
    list.innerHTML = "";
  }
}

function addCategories() {
  const categories = [...new Set(state.events.map(event => event.category))].sort();
  categories.forEach(category => {
    const option = document.createElement("option");
    option.value = category;
    option.textContent = category;
    categoryFilter.appendChild(option);
  });
}

function applyFilters() {
  const term = searchInput.value.trim().toLowerCase();
  const category = categoryFilter.value;
  const direction = sortSelect.value === "asc" ? 1 : -1;

  state.filtered = state.events
    .filter(event => {
      const searchable = `${event.title} ${event.venue}`.toLowerCase();
      return searchable.includes(term) && (category === "all" || event.category === category);
    })
    .sort((a, b) => (new Date(a.date) - new Date(b.date)) * direction);

  state.page = 1;
  renderEvents();
  renderPagination();
}

function renderEvents() {
  const start = (state.page - 1) * state.pageSize;
  const pageItems = state.filtered.slice(start, start + state.pageSize);
  status.textContent = `${state.filtered.length} event(s) found`;
  list.innerHTML = pageItems.length ? pageItems.map(event => `
    <article class="event-card">
      <h2>${event.title}</h2>
      <p><strong>Date:</strong> ${formatDate(event.date)}</p>
      <p><strong>Venue:</strong> ${event.venue}</p>
      <p><strong>Category:</strong> ${event.category}</p>
    </article>
  `).join("") : "<p>No matching events found.</p>";
}

function renderPagination() {
  const totalPages = Math.ceil(state.filtered.length / state.pageSize);
  pagination.innerHTML = "";
  for (let page = 1; page <= totalPages; page += 1) {
    const button = document.createElement("button");
    button.textContent = page;
    button.setAttribute("aria-label", `Go to page ${page}`);
    button.className = page === state.page ? "active" : "";
    button.addEventListener("click", () => {
      state.page = page;
      renderEvents();
      renderPagination();
    });
    pagination.appendChild(button);
  }
}

function formatDate(value) {
  return new Date(value).toLocaleDateString("en-IN", {
    day: "2-digit", month: "short", year: "numeric"
  });
}

searchInput.addEventListener("input", applyFilters);
categoryFilter.addEventListener("change", applyFilters);
sortSelect.addEventListener("change", applyFilters);
loadEvents();'''


CSS = '''* { box-sizing: border-box; }
body { margin: 0; font-family: Arial, sans-serif; background: #0b162c; color: #f5f7fb; }
.container { width: min(100% - 32px, 1000px); margin: 32px auto; }
h1 { color: #ec9f05; }
.controls { display: grid; grid-template-columns: auto 1fr auto 1fr auto 1fr; gap: 10px; align-items: center; padding: 16px; background: #152644; border-radius: 10px; }
input, select { width: 100%; padding: 10px; border: 1px solid #7184a8; border-radius: 6px; background: #ffffff; color: #17233b; }
.status { color: #b7c5dd; }
.event-grid { display: grid; grid-template-columns: repeat(2, 1fr); gap: 16px; }
.event-card { padding: 18px; border: 1px solid #52698e; border-radius: 10px; background: #13213a; }
.event-card h2 { margin-top: 0; color: #ec9f05; font-size: 20px; }
.pagination { display: flex; gap: 8px; margin-top: 20px; }
.pagination button { padding: 8px 12px; border: 0; border-radius: 5px; cursor: pointer; }
.pagination .active { background: #ec9f05; color: #17233b; font-weight: bold; }
@media (max-width: 720px) { .controls, .event-grid { grid-template-columns: 1fr; } }
'''


def main():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    title = doc.add_paragraph(style="Title")
    para(title, before=0, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    title_run = title.add_run("Practical 6 Fetch API Search and Filter")
    font(title_run, size=24, bold=True, color="000000")
    sub = doc.add_paragraph()
    para(sub, after=14, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = sub.add_run("Student Hub Portal")
    font(r, size=13, color="2E74B5")

    heading(doc, "Question")
    labeled(doc, "Practical 6", "Rendering External JSON Data using Fetch API, Search and Filter")
    labeled(doc, "Problem Definition", "Display event lists, student profiles, notices, or FAQs from an external JSON file using the Fetch API. Implement dynamic rendering, search, filter, sorting, and pagination.")

    heading(doc, "Key Questions and Answers", 2)
    questions = [
        ("How is JSON fetched, parsed, and rendered?", "The fetch function requests events.json and returns a response. The response is converted to JavaScript objects with response.json(), and the objects are inserted into the page as event cards."),
        ("Which array methods are used for search, filter, sort, and pagination?", "filter() selects matching events, sort() orders events by date, map() creates display markup, and slice() selects the records for the current page."),
        ("How are loading and error states handled?", "The page initially shows a loading message. If the request fails or the response is not successful, catch() displays an error message and clears the event list."),
        ("How is modularity maintained in JavaScript files?", "The code is divided into small functions such as loadEvents, addCategories, applyFilters, renderEvents, renderPagination, and formatDate. Each function has one clear responsibility."),
    ]
    for q, a in questions:
        labeled(doc, "Q", q)
        labeled(doc, "Answer", a)

    heading(doc, "Analysis", 2)
    labeled(doc, "A", "The application loads event records from events.json instead of writing every event directly in HTML. Search, category filtering, date sorting, and pagination are applied to the JavaScript array before the visible cards are updated.")
    heading(doc, "Interpretation", 2)
    labeled(doc, "I", "This approach is useful for modern web applications because the same page can display changing data without rewriting the HTML. In a real project, the JSON file can later be replaced by an API endpoint.")

    heading(doc, "Code Required", 2)
    labeled(doc, "Testing Note", "Run the files through a local server such as Live Server. Fetch may be blocked when the HTML file is opened directly with a file path in the browser.")
    code_block(doc, "HTML Code - events-fetch.html", HTML)
    code_block(doc, "JSON Code - events.json", JSON_DATA)
    code_block(doc, "JavaScript Code - practical6.js", JS)
    code_block(doc, "CSS Code - practical6.css", CSS)

    doc.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
