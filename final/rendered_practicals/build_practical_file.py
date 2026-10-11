from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DOCX = ROOT / "Student_Hub_Practical_File_All_HTML_CSS.docx"


def read_text(name: str) -> str:
    return (ROOT / name).read_text(encoding="utf-8")


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def set_paragraph(paragraph, before=0, after=0, line=1.15, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_title_page(doc: Document):
    p = doc.add_paragraph()
    set_paragraph(p, before=120, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Student Hub Portal")
    set_run_font(r, name="Calibri", size=24, bold=True, color="1F1F1F")

    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Practical File")
    set_run_font(r, name="Calibri", size=18, bold=True, color="2E74B5")

    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=16, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Complete HTML and CSS Project File")
    set_run_font(r, name="Calibri", size=12.5, color="4D4D4D")

    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=16, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("All HTML pages and the common stylesheet have been arranged practical-wise.")
    set_run_font(r, name="Calibri", size=11, color="4D4D4D")

    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Student Hub Portal Source Code")
    set_run_font(r, name="Calibri", size=10.5, italic=True, color="666666")

    doc.add_page_break()


def add_heading_text(doc: Document, text: str):
    p = doc.add_paragraph()
    set_paragraph(p, before=6, after=6, line=1.1)
    r = p.add_run(text)
    set_run_font(r, name="Calibri", size=15, bold=True, color="2E74B5")


def add_question(doc: Document, text: str):
    p = doc.add_paragraph()
    set_paragraph(p, before=0, after=4, line=1.15)
    r1 = p.add_run("Question: ")
    set_run_font(r1, name="Calibri", size=11, bold=True, color="000000")
    r2 = p.add_run(text)
    set_run_font(r2, name="Calibri", size=11, color="000000")


def add_label(doc: Document, text: str):
    p = doc.add_paragraph()
    set_paragraph(p, before=6, after=3, line=1.0)
    r = p.add_run(text)
    set_run_font(r, name="Calibri", size=11, bold=True, color="1F4D78")


def add_code_block(doc: Document, code: str):
    lines = code.rstrip("\n").splitlines()
    for line in lines:
        p = doc.add_paragraph()
        set_paragraph(p, before=0, after=0, line=1.0)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.05)
        r = p.add_run(line.rstrip())
        set_run_font(r, name="Courier New", size=8.3, color="000000")


def add_practical(doc: Document, no: int, title: str, question: str, sources):
    add_heading_text(doc, f"Practical No. {no} - {title}")
    add_question(doc, question)
    for label, code in sources:
        add_label(doc, label)
        add_code_block(doc, code)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)


def practical_data():
    return [
        (
            1,
            "Home Page (index.html)",
            "Create the home page of the Student Hub Portal using HTML and an external CSS file so the page shows a header, navigation bar, hero area, notice section, quick links, and footer.",
            [
                ("HTML Code - index.html", read_text("index.html")),
                ("CSS Code - style.css", read_text("style.css")),
            ],
        ),
        (
            2,
            "About Page (about.html)",
            "Create the About page to present the mission and vision of Student Hub Portal in a simple, readable layout.",
            [
                ("HTML Code - about.html", read_text("about.html")),
            ],
        ),
        (
            3,
            "Events Page (events.html)",
            "Create the Events page so students can search events, view event details, and register from a clean table layout.",
            [
                ("HTML Code - events.html", read_text("events.html")),
            ],
        ),
        (
            4,
            "FAQ Page (faq.html)",
            "Create the FAQ page and answer the common questions students may ask about registration, login, events, and contact options.",
            [
                ("HTML Code - faq.html", read_text("faq.html")),
            ],
        ),
        (
            5,
            "Contact & Feedback Page (contact.html)",
            "Create the Contact and Feedback page so users can find contact details and submit feedback through a form.",
            [
                ("HTML Code - contact.html", read_text("contact.html")),
            ],
        ),
        (
            6,
            "Login Page (login.html)",
            "Create a login page for the Student Hub Portal with fields for email and password, a remember-me option, a login button, and helpful navigation links.",
            [
                ("HTML Code - login.html", read_text("login.html")),
            ],
        ),
        (
            7,
            "Student Registration Page (register.html)",
            "Create a student registration form that collects personal details, course information, gender, password confirmation, and terms acceptance before submission.",
            [
                ("HTML Code - register.html", read_text("register.html")),
            ],
        ),
        (
            8,
            "Dashboard Page (dashboard.html)",
            "Create the student dashboard page to display menu options, summary cards, upcoming events, and important notices after login.",
            [
                ("HTML Code - dashboard.html", read_text("dashboard.html")),
            ],
        ),
        (
            9,
            "Profile Page (profile.html)",
            "Create the profile page to show the student photo area, personal details, and an Edit Profile button in a centered layout.",
            [
                ("HTML Code - profile.html", read_text("profile.html")),
            ],
        ),
        (
            10,
            "Admin Panel (admin.html)",
            "Create the admin panel page to show summary counters, recent activities, and navigation for administration tasks.",
            [
                ("HTML Code - admin.html", read_text("admin.html")),
            ],
        ),
        (
            11,
            "Common Stylesheet (style.css)",
            "Create the external stylesheet used across the Student Hub Portal pages so the navigation, buttons, cards, forms, and footer share one consistent visual style.",
            [
                ("CSS Code - style.css", read_text("style.css")),
            ],
        ),
    ]


def main():
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    data = practical_data()
    for index, item in enumerate(data):
        if index:
            doc.add_page_break()
        add_practical(doc, *item)

    doc.save(OUT_DOCX)
    print(f"Saved {OUT_DOCX}")


if __name__ == "__main__":
    main()
