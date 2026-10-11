from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT_DOCX = ROOT / "Student_Hub_Practical_File_Upto_5.docx"


def read_text(name: str) -> str:
    return (ROOT / name).read_text(encoding="utf-8")


def set_run_font(run, name="Calibri", size=11, bold=None, italic=None, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=0, line=1.15, align=None):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    if align is not None:
        paragraph.alignment = align


def add_title_page(doc: Document):
    p = doc.add_paragraph()
    style_paragraph(p, before=120, after=8, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Student Hub Portal")
    set_run_font(r, size=24, bold=True, color="1F1F1F")

    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=6, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Practical File")
    set_run_font(r, size=18, bold=True, color="2E74B5")

    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=14, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Practicals 1 to 5 - HTML, CSS and JavaScript")
    set_run_font(r, size=12.5, color="4D4D4D")

    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=16, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Practical-wise write-up with questions, answers, analysis, interpretation, and source code.")
    set_run_font(r, size=11, color="4D4D4D")

    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=0, line=1.0, align=WD_ALIGN_PARAGRAPH.CENTER)
    r = p.add_run("Student Hub Portal Source Code")
    set_run_font(r, size=10.5, italic=True, color="666666")

    doc.add_page_break()


def add_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    style_paragraph(p, before=6, after=6, line=1.1)
    r = p.add_run(text)
    set_run_font(r, size=15, bold=True, color="2E74B5")


def add_subheading(doc: Document, text: str):
    p = doc.add_paragraph()
    style_paragraph(p, before=4, after=2, line=1.05)
    r = p.add_run(text)
    set_run_font(r, size=11.5, bold=True, color="1F4D78")


def add_body(doc: Document, label: str, text: str):
    p = doc.add_paragraph()
    style_paragraph(p, before=0, after=3, line=1.15)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=11, bold=True, color="000000")
    r2 = p.add_run(text)
    set_run_font(r2, size=11, color="000000")


def add_code_heading(doc: Document, text: str):
    p = doc.add_paragraph()
    style_paragraph(p, before=6, after=3, line=1.0)
    r = p.add_run(text)
    set_run_font(r, size=11, bold=True, color="1F4D78")


def add_code_block(doc: Document, code: str):
    for line in code.rstrip("\n").splitlines():
        p = doc.add_paragraph()
        style_paragraph(p, before=0, after=0, line=1.0)
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.05)
        r = p.add_run(line.rstrip())
        stripped = line.strip()
        if stripped.startswith("<!--") or stripped.startswith("//") or stripped.startswith("/*") or stripped.startswith("*"):
            color = "008000"
        elif stripped.startswith("<"):
            color = "0000CC"
        elif ":" in stripped and stripped.endswith(";"):
            color = "006666"
        elif any(word in stripped.split() for word in ("const", "let", "function", "return", "if", "else")):
            color = "7F0055"
        else:
            color = "333333"
        set_run_font(r, name="Courier New", size=9, color=color)


def add_practical(doc: Document, no: int, title: str, question: str, key_questions, analysis, interpretation, sources):
    add_heading(doc, f"Practical No. {no} - {title}")
    add_body(doc, "Question", question)

    add_subheading(doc, "Key Questions")
    for item in key_questions:
        add_body(doc, "Q", item["question"])
        add_body(doc, "Answer", item["answer"])

    add_subheading(doc, "Analysis")
    add_body(doc, "A", analysis)

    add_subheading(doc, "Interpretation")
    add_body(doc, "I", interpretation)

    for label, code in sources:
        add_code_heading(doc, label)
        add_code_block(doc, code)


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.45)
    section.footer_distance = Inches(0.45)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)


def practical_data():
    style_css = read_text("style.css")
    return [
        (
            1,
            "Home Page (index.html)",
            "Create the home page of the Student Hub Portal using HTML and an external CSS file so the page shows a header, navigation bar, hero area, notice section, quick links, and footer.",
            [
                {
                    "question": "How does the header help the user identify the portal quickly?",
                    "answer": "The header shows the logo and portal name at the top, so the user immediately knows which site they are viewing.",
                },
                {
                    "question": "What navigation links make the home page easy to explore?",
                    "answer": "The menu links to Home, About, Events, FAQ, Contact, and Feedback, which gives direct access to the main sections of the portal.",
                },
                {
                    "question": "How do the notice, events, and quick links sections organize the content?",
                    "answer": "They split the main content into three clear cards, so updates, upcoming activities, and shortcuts are easy to read and find.",
                },
            ],
            "This page uses a clear visual hierarchy: a branded header, centered hero message, three content cards, and a footer. The external stylesheet keeps the layout, colors, buttons, and spacing consistent.",
            "The home page acts as the landing page for the portal. It should give a fast overview and direct the user toward login, registration, events, and other major sections.",
            [
                ("HTML Code - index.html", read_text("index.html")),
                ("CSS Code - style.css", style_css),
            ],
        ),
        (
            2,
            "About Page (about.html)",
            "Create the About page to present the mission and vision of Student Hub Portal in a simple, readable layout.",
            [
                {
                    "question": "What is the main purpose of the About page?",
                    "answer": "The About page explains what the portal is for and gives the user a short introduction to the project.",
                },
                {
                    "question": "How are the mission and vision separated for clarity?",
                    "answer": "They are placed in separate boxes, which keeps each idea distinct and easy to compare.",
                },
                {
                    "question": "Why is a two-card layout useful for this content?",
                    "answer": "A two-card layout keeps the page balanced and lets the reader focus on one message at a time.",
                },
            ],
            "The About page is written in a simple card layout so the mission and vision can be read independently. The structure is short, direct, and easy to scan.",
            "The page communicates what the portal stands for and what it aims to achieve. It supports trust and gives the project a clear identity.",
            [("HTML Code - about.html", read_text("about.html"))],
        ),
        (
            3,
            "Events Page (events.html)",
            "Create the Events page so students can search events, view event details, and register from a clean table layout.",
            [
                {
                    "question": "What information does each event card provide?",
                    "answer": "Each event card gives the event name, date, venue, category, and a register button, so the user sees all important details at once.",
                },
                {
                    "question": "How does the search box support event discovery?",
                    "answer": "It lets the student type a keyword and quickly look for a specific event instead of scanning the full list manually.",
                },
                {
                    "question": "Why is a table layout suitable for this page?",
                    "answer": "A table keeps the event details aligned in rows and columns, which makes the page neat and easy to compare.",
                },
            ],
            "The Events page combines a search field with a structured event listing. Each row gives the event title, date, venue, category, and a register button for quick action.",
            "This page is meant for browsing and action. The design helps users find an event and register without needing to open multiple pages.",
            [("HTML Code - events.html", read_text("events.html"))],
        ),
        (
            4,
            "FAQ Page (faq.html)",
            "Create the FAQ page and answer the common questions students may ask about registration, login, events, and contact options.",
            [
                {
                    "question": "Which student problems are answered on this page?",
                    "answer": "The page answers common doubts about registration, login, joining events, and contacting the admin team.",
                },
                {
                    "question": "How does the FAQ reduce the need for support messages?",
                    "answer": "It gives simple answers to repeated questions, so students can solve basic issues without sending a new query every time.",
                },
                {
                    "question": "Why are the questions written in a direct format?",
                    "answer": "Direct questions match the way students usually ask for help, so the answers are easier to understand quickly.",
                },
            ],
            "The FAQ page uses repeated question-answer blocks separated by horizontal lines. This gives the page a simple reading path and keeps the answers easy to locate.",
            "The page works as a support section. It helps students solve common doubts on their own and keeps the portal easier to use.",
            [("HTML Code - faq.html", read_text("faq.html"))],
        ),
        (
            5,
            "Contact & Feedback Page (contact.html)",
            "Create the Contact and Feedback page so users can find contact details and submit feedback through a form.",
            [
                {
                    "question": "What contact details are shown first on the page?",
                    "answer": "The page shows the email address, phone number, and campus address first so the user can contact support quickly.",
                },
                {
                    "question": "How does the feedback form encourage user response?",
                    "answer": "The form is short and simple, so the user can send a message without spending too much time.",
                },
                {
                    "question": "Why are contact details and feedback kept in separate sections?",
                    "answer": "Separating them keeps the page organized and makes it clear whether the user is reading information or submitting a response.",
                },
            ],
            "The page is split into a contact information card and a feedback form card. This separation helps the user either read details quickly or take action by sending a message.",
            "The page serves both communication and collection purposes. It gives the user a way to reach support and to share suggestions in a structured form.",
            [("HTML Code - contact.html", read_text("contact.html"))],
        ),
        (
            6,
            "Login Page (login.html)",
            "Create a login page for the Student Hub Portal with fields for email and password, a remember-me option, a login button, and helpful navigation links.",
            [
                {
                    "question": "What input fields are required to sign in?",
                    "answer": "The user must enter an email address and password before submitting the login form.",
                },
                {
                    "question": "How does the remember-me option improve usability?",
                    "answer": "It allows the user to stay signed in on the device, which makes future logins faster and easier.",
                },
                {
                    "question": "Why is there a forgot-password link below the form?",
                    "answer": "The link helps users recover access if they forget their password, which makes the page more practical.",
                },
            ],
            "The login form is kept compact so the user can enter credentials quickly. The checkbox and support links make the page feel complete without becoming crowded.",
            "The page focuses on authentication. It gives users a direct path into the dashboard while keeping the form simple and familiar.",
            [("HTML Code - login.html", read_text("login.html"))],
        ),
        (
            7,
            "Student Registration Page (register.html)",
            "Create a student registration form that collects personal details, course information, gender, password confirmation, and terms acceptance before submission.",
            [
                {
                    "question": "Which student details are collected in the form?",
                    "answer": "The form collects name, email, mobile number, course, semester, gender, password, and confirm password.",
                },
                {
                    "question": "Why are password and confirm-password fields both included?",
                    "answer": "Both fields are included to reduce typing mistakes and make sure the student enters the correct password.",
                },
                {
                    "question": "How does the terms checkbox make the form more complete?",
                    "answer": "The checkbox confirms that the user agrees to the terms and conditions before creating the account.",
                },
            ],
            "The registration page contains a larger form because it collects several student details. The fields are arranged in a straight order so the form remains easy to fill in.",
            "This page captures the information needed to create a student account. The design supports accuracy by asking for all required details before submission.",
            [("HTML Code - register.html", read_text("register.html"))],
        ),
        (
            8,
            "Dashboard Page (dashboard.html)",
            "Create the student dashboard page to display menu options, summary cards, upcoming events, and important notices after login.",
            [
                {
                    "question": "What summary information is shown on the dashboard?",
                    "answer": "The dashboard shows the total events count, registered count, upcoming event, and latest notice.",
                },
                {
                    "question": "How does the sidebar menu help navigation?",
                    "answer": "The sidebar gives quick buttons for Home, Profile, Events, Notices, and Logout, so movement inside the portal is simple.",
                },
                {
                    "question": "Why is the dashboard split into multiple boxes?",
                    "answer": "Separate boxes make the information easier to scan and help the important items stand out clearly.",
                },
            ],
            "The dashboard mixes navigation and live summary blocks. It is designed like a control panel so the user can move around the portal and see important updates at a glance.",
            "The dashboard is the main post-login workspace. It gives the student a fast view of the most important information and action points.",
            [("HTML Code - dashboard.html", read_text("dashboard.html"))],
        ),
        (
            9,
            "Profile Page (profile.html)",
            "Create the profile page to show the student photo area, personal details, and an Edit Profile button in a centered layout.",
            [
                {
                    "question": "What student data is shown on the profile page?",
                    "answer": "The page shows the student's name, email, mobile number, course, and semester.",
                },
                {
                    "question": "How does the profile picture placeholder support the layout?",
                    "answer": "It gives a visible profile area even when no image is uploaded, so the page still looks complete.",
                },
                {
                    "question": "Why is an edit button useful on this page?",
                    "answer": "The edit button lets the user update profile information, which makes the page more useful and interactive.",
                },
            ],
            "The profile page presents a centered information card with personal details and a clear action button. The layout is simple, which makes the profile feel organized and readable.",
            "The page represents the student identity inside the portal. It provides a quick summary of the user and allows profile updates when needed.",
            [("HTML Code - profile.html", read_text("profile.html"))],
        ),
        (
            10,
            "Admin Panel (admin.html)",
            "Create the admin panel page to show summary counters, recent activities, and navigation for administration tasks.",
            [
                {
                    "question": "Which portal statistics are shown for the admin?",
                    "answer": "The panel shows the number of students, events, feedback entries, and notices.",
                },
                {
                    "question": "How do summary counters help with monitoring?",
                    "answer": "They give a quick numerical view of portal activity, which helps the admin understand the current status at a glance.",
                },
                {
                    "question": "Why is a recent activity list useful for administration?",
                    "answer": "It highlights the latest actions on the portal, so the admin can track changes and activity easily.",
                },
            ],
            "The admin panel uses compact summary boxes and a simple activity list. This helps the administrator track the portal status without needing a complex interface.",
            "The page is intended for monitoring and control. It gives a quick operational overview and supports administrative decision-making.",
            [("HTML Code - admin.html", read_text("admin.html"))],
        ),
        (
            11,
            "Common Stylesheet (style.css)",
            "Create the external stylesheet used across the Student Hub Portal pages so the navigation, buttons, cards, forms, and footer share one consistent visual style.",
            [
                {
                    "question": "Which shared visual elements are controlled by the stylesheet?",
                    "answer": "The stylesheet controls the body colors, header, navigation, main cards, buttons, forms, links, and footer styles.",
                },
                {
                    "question": "How does the stylesheet keep the site design consistent?",
                    "answer": "It applies the same fonts, colors, spacing, and button style across all pages, so the project looks unified.",
                },
                {
                    "question": "Why is using one CSS file better than repeating styles in each page?",
                    "answer": "One CSS file is easier to manage, reduces duplication, and makes future edits much faster.",
                },
            ],
            "The stylesheet defines the site color palette, layout spacing, button behavior, card styling, and form appearance. It gives all pages the same visual identity.",
            "The CSS file is the design backbone of the project. It reduces repetition, keeps styling consistent, and makes future changes easier to manage.",
            [("CSS Code - style.css", style_css)],
        ),
    ]


def practical_data_upto_5():
    style_css = read_text("style.css")
    script_js = read_text("script.js")
    html_files = [
        "index.html", "about.html", "events.html", "faq.html", "contact.html",
        "login.html", "register.html", "dashboard.html", "profile.html", "admin.html",
    ]
    semantic_code = [(f"HTML Code - {name}", read_text(name)) for name in html_files]
    return [
        (
            1,
            "Project Initiation, Requirement Analysis, Sitemap, Wireframe, and GitHub Setup",
            "Initiate the semester-long StudentHub portal by identifying the problem scope, user roles, key modules, navigation flow, and minimum 10 pages. Create a sitemap, low-fidelity wireframe, project folder structure, README file, and GitHub repository.",
            [
                {"question": "What is URL and parts of URL?", "answer": "A URL is the address used to open a resource on the internet. Its main parts are the protocol such as https, domain name, path, query string, and fragment."},
                {"question": "How does HTML file processed in web browser?", "answer": "The browser reads the HTML file, builds the DOM tree, loads linked CSS and JavaScript files, and then displays the final page on the screen."},
                {"question": "How will page navigation flow managed among all HTML pages?", "answer": "Navigation is managed with anchor tags and relative links such as index.html, about.html, events.html, and register.html. A common navigation bar keeps the flow consistent."},
                {"question": "How will GitHub commits be maintained after each practical?", "answer": "After completing and testing each practical, the changed files should be added, committed with a clear message, and pushed to the GitHub repository."},
            ],
            "The Student Hub project is divided into clear pages for home, information, events, support, authentication, student details, and administration. Separate HTML, CSS, and JavaScript files keep the work organized.",
            "A sitemap and wireframe give the project a clear direction before coding starts. This reduces confusion about page connections and makes later GitHub commits easier to understand.",
            [("Project File Structure", "index.html\nabout.html\nevents.html\nfaq.html\ncontact.html\nlogin.html\nregister.html\ndashboard.html\nprofile.html\nadmin.html\nstyle.css\nscript.js")],
        ),
        (
            2,
            "Semantic HTML5 Pages with Accessibility-Ready Structure",
            "Develop static HTML5 skeletons for at least 10 StudentHub pages such as Home, About, Register, Login, Dashboard, Events, Profile, Contact, Admin, FAQ, and Feedback. Use semantic tags and accessibility-friendly structure.",
            [
                {"question": "Are semantic tags such as header, nav, main, section, article, aside, and footer used properly?", "answer": "Yes. The pages use header for branding, nav for links, main for page content, sections or cards for grouped information, and footer for common bottom links."},
                {"question": "Are form labels, headings, image alt text, and navigation links accessible?", "answer": "The form fields have labels, the pages use headings, the logo has alt text, and navigation is provided through normal links. These features help keyboard and screen-reader users understand the page."},
                {"question": "Is the page structure consistent across all pages?", "answer": "Yes. The pages follow a common pattern with a header, navigation menu, main content area, footer, and linked CSS and JavaScript files."},
            ],
            "The project has ten linked HTML pages with a common layout. Semantic elements make the structure easier for browsers, developers, and assistive technologies to understand.",
            "The pages form a consistent static website foundation. A user can move from the home page to information, events, registration, login, and support pages without losing the common site structure.",
            semantic_code,
        ),
        (
            3,
            "Responsive UI Design using CSS Grid, Flexbox, and Bootstrap/Tailwind",
            "Design responsive layouts for Home, About, Registration, Dashboard, and Events pages using CSS Grid and Flexbox. Ensure mobile-first responsiveness.",
            [
                {"question": "How does the layout adapt for mobile, tablet, and desktop screens?", "answer": "Flexbox allows items to wrap and share available space, while media queries change widths and the navigation layout for smaller screens. The same content can fit different screen sizes."},
                {"question": "Which layout technique is used and why?", "answer": "Flexbox is used for horizontal groups such as the header, navigation, form grid, and cards. It is suitable because it aligns items and distributes space with less code."},
                {"question": "Are typography, spacing, colors, and contrast readable?", "answer": "The stylesheet uses a dark background, light text, contrasting yellow buttons, consistent padding, and readable font sizes. These choices make the important content visible."},
                {"question": "Is the framework used consistently without unnecessary inline CSS?", "answer": "The project mainly uses one external style.css file. Some page-specific inline styles are present, but the repeated site-wide styles are kept in the common stylesheet."},
            ],
            "The CSS file defines the common color palette, cards, buttons, forms, navigation, and responsive rules. Flexbox is used for groups of elements and media queries adjust the layout for smaller screens.",
            "The interface is designed to remain usable on desktop and mobile screens. Centralizing the main styles makes the pages look like one portal instead of separate unrelated pages.",
            [("CSS Code - style.css", style_css)],
        ),
        (
            4,
            "JavaScript DOM Manipulation, Event Handling, and UI Interactivity",
            "Add dynamic UI components such as collapsible FAQ, modal popup, image/content slider, notification banner, hamburger menu, and light/dark theme switcher using JavaScript.",
            [
                {"question": "How are DOM elements selected and modified?", "answer": "The script uses getElementById, querySelector, and querySelectorAll. It changes class names, text, display styles, and the current slide to update the page."},
                {"question": "Are event listeners attached correctly?", "answer": "The HTML buttons call functions for clicks, and a DOMContentLoaded listener starts the first slider. A window click handler also closes a modal when the user clicks outside it."},
                {"question": "Is localStorage used for remembering theme choice?", "answer": "Yes. The selected theme is stored in localStorage, and the saved value is checked when the page loads so the preference can be restored."},
                {"question": "Does interactivity improve usability without breaking accessibility?", "answer": "The features make the page easier to use by hiding repeated information, showing announcements, and supporting mobile navigation. Clear buttons and visible text keep the controls understandable."},
            ],
            "The JavaScript file contains small functions for theme switching, menu display, banner closing, FAQ expansion, modal display, and slider movement. Each function changes only the DOM elements needed for that feature.",
            "The portal becomes interactive without a server. Users can control the menu, announcements, FAQ answers, theme, and slider directly in the browser.",
            [("JavaScript Code - script.js", script_js)],
        ),
        (
            5,
            "Registration Form with Frontend Validation and User-Friendly Error Handling",
            "Create a student registration form with HTML5 input types and JavaScript validation for name, email, mobile number, password, confirm password, course, year, gender, and terms acceptance. Use Regular Expression for validations.",
            [
                {"question": "Are correct input types and attributes used?", "answer": "Yes. The form uses text, email, tel, number, password, select, radio, and checkbox controls. Labels, placeholders, min and max values, and validation checks support correct input."},
                {"question": "Are validation errors displayed near the relevant fields?", "answer": "Yes. Each important field has a nearby error element such as nameError or emailError. The validation function writes the message there and adds an input-error class."},
                {"question": "Is password strength checked?", "answer": "Yes. A regular expression checks that the password has at least six characters, at least one letter, and at least one number. The confirm password is also compared with it."},
                {"question": "Is the form accessible using keyboard and screen-reader-friendly labels?", "answer": "The controls are connected to visible labels using for and id values, and the radio buttons have a shared name. This gives keyboard and assistive-technology users understandable fields."},
            ],
            "The registration form collects the required student details and validates them before accepting the form. Invalid fields receive a message and an error class, while valid fields receive a success class.",
            "Client-side validation gives immediate feedback and prevents incomplete or incorrectly formatted data from being submitted. Server-side validation would still be needed in a real application.",
            [("HTML Code - register.html", read_text("register.html")), ("JavaScript Code - validation in script.js", script_js)],
        ),
    ]


def main():
    doc = Document()
    configure_document(doc)
    add_title_page(doc)

    for i, item in enumerate(practical_data_upto_5()):
        if i:
            doc.add_page_break()
        add_practical(doc, *item)

    doc.save(OUT_DOCX)
    print(f"Saved {OUT_DOCX}")


if __name__ == "__main__":
    main()
